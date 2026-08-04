"""workbook_cells.py — cell-preview payloads for the workbook citation tier.

PM-review Sprint 2 Task 4 (spec §P1-9). The workbook drawer used to say
"Sheet: Exhibit P-1 / Cells: O839 O840 O841" and offer a 4 MB .xlsx download.
The product promise is that a citation opens the source with the exact cell
highlighted, so the drawer now renders a small HTML table of the cited cells
plus surrounding rows — and, when a fact sums several cells, the arithmetic
that joins them.

That table is built HERE, at export time, from the same workbook files the
loaders read (govbudget.jbooks.p1_loader / rollup_loader) and the 5B-1 gate
re-derives against (verify_phase5b1._verify_workbook). Nothing is inferred:
every value, label and column header is read out of the workbook.

Payload, one per workbook citation fact_id::

    {
      "sheet":      "Exhibit P-1",
      "col":        "O",                        # the cited column's letters
      "col_header": "FY 2024 Actuals Amount",   # that column's header text
      "units":      "USD thousands",
      "total":      5565655.0,                  # == citation amount_thousands
      "rows": [
        {"r": 837, "code": "B02100", "title": "B-21 Raider",
         "note": "C (FY 2025 for FY 2026) (M)", "flag": "Non-Add"},
        {"r": 839, "code": "ATA000", "title": "F-35",
         "note": "Weapon System Cost", "flag": "Add",
         "v": 5493772.0, "cited": true},
        ...
      ]
    }

Rows are in workbook row order; `cited` marks the citation's own cells;
absent optional keys mean the workbook has no such column / the cell is blank
(never a fabricated zero).

SIZE: these payloads ship as their OWN sharded sidecar
(json/workbook-cells/{fact_id[:2]}.json), NOT inside the citation rows.
Folding them into citations.json would inflate every embedded per-page
citation slice — i.e. the static HTML of every figure-bearing page — and the
cite-shards by ~25%. The drawer fetches one shard the first time a workbook
citation is opened. See lib/workbook-cells.ts for the client half.

LOUDNESS: a preview that cannot be built exactly — missing sheet, out-of-range
cell, blank cited cell, or a cited-cell sum that disagrees with the citation's
amount — raises. A workbook drawer showing cells that do not add up to the
figure above them would be worse than the defect it replaces.

The ONE tolerated case is a workbook FILE that was never exported. That is not
a preview defect, it is a missing document, and it already has a gate that
names it: verify_phase5b1._verify_workbook fails with "workbook not found" for
every workbook citation whose .xlsx is absent. Callers opt into tolerance by
passing a `missing_workbooks` set to collect the shas (export_site does, and
prints the count); pass nothing and a missing file raises. Either way the
citation ships previewless, and gate 4 leg 3a — "every workbook citation has a
cell preview" — fails the build. Two gates, no silent gap.
"""
from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

# Rows of context kept above and below the cited block (§P1-9: "the cited
# cells plus ~2 surrounding rows"). Clamped to the workbook body — the
# preview never walks above the header row or past the last row.
CONTEXT_ROWS = 2

_CELL_RE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")

# Header-row detection: the display workbooks all carry these two, and the
# loaders use the same probe (first 20 rows).
_HEADER_REQUIRED = ("Account", "Organization")
_HEADER_PROBE_ROWS = 20

# Label columns, by normalized header name, most specific first. Every P-1 /
# R-1 / P-1R edition from PB2017 on is covered by one alternative per slot.
_CODE_HEADERS = ("Budget Line Item", "PE/BLI", "Line Item", "Line Number")
_TITLE_HEADERS = (
    "Budget Line Item (BLI) Title",
    "Program Element/Budget Line Item (BLI) Title",
    "Line Item Title",
)
_NOTE_HEADERS = ("Cost Type Title", "Budget SubActivity (BSA) Title")
_FLAG_HEADERS = ("Add/Non-Add",)


def norm_header(v) -> str:
    """Normalize a header cell.

    Mirrors govbudget.jbooks.rollup_loader.norm_header: PB2017–PB2023 display
    workbooks wrap headers across lines ('Budget\\nActivity') and space the
    slash ('PE / BLI'); PB2024+ headers are single-line and pass through.
    """
    h = re.sub(r"\s+", " ", str(v)).strip()
    return re.sub(r"\s*/\s*", "/", h)


def parse_cell_ref(ref: str) -> tuple[str, int]:
    """'O839' -> ('O', 839). Raises ValueError on anything else."""
    m = _CELL_RE.match(str(ref).strip().upper())
    if not m:
        raise ValueError(f"not a cell reference: {ref!r}")
    return m.group(1), int(m.group(2))


def column_letters_to_index(letters: str) -> int:
    """'A' -> 1, 'O' -> 15, 'AA' -> 27 (1-based, like openpyxl)."""
    idx = 0
    for ch in letters.upper():
        if not ("A" <= ch <= "Z"):
            raise ValueError(f"not column letters: {letters!r}")
        idx = idx * 26 + (ord(ch) - 64)
    if idx <= 0:
        raise ValueError(f"not column letters: {letters!r}")
    return idx


def _cell_str(v) -> str | None:
    """Trimmed text of a workbook cell, or None when it is empty."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _find_header_row(rows: list[tuple]) -> tuple[int, dict[int, str]]:
    """(1-based header row index, {0-based column -> normalized header})."""
    for i, row in enumerate(rows[:_HEADER_PROBE_ROWS], start=1):
        headers = {j: norm_header(v) for j, v in enumerate(row) if v is not None}
        vals = set(headers.values())
        if all(h in vals for h in _HEADER_REQUIRED):
            return i, headers
    raise ValueError(
        f"no header row with {_HEADER_REQUIRED} in the first"
        f" {_HEADER_PROBE_ROWS} rows"
    )


def _pick_column(headers: dict[int, str], names: tuple[str, ...]) -> int | None:
    for name in names:
        for j, h in headers.items():
            if h == name:
                return j
    return None


class _Sheet:
    """One workbook sheet, materialized once and shared by every citation."""

    def __init__(self, name: str, rows: list[tuple]):
        self.name = name
        self.rows = rows
        self.header_row, self.headers = _find_header_row(rows)
        self.code_col = _pick_column(self.headers, _CODE_HEADERS)
        self.title_col = _pick_column(self.headers, _TITLE_HEADERS)
        self.note_col = _pick_column(self.headers, _NOTE_HEADERS)
        self.flag_col = _pick_column(self.headers, _FLAG_HEADERS)

    def value(self, row_1based: int, col_0based: int):
        row = self.rows[row_1based - 1]
        return row[col_0based] if col_0based < len(row) else None

    def label_row(self, row_1based: int, col_0based: int) -> dict:
        """The preview row for one workbook row, minus the `cited` marker."""
        out: dict = {"r": row_1based}
        for key, col in (
            ("code", self.code_col),
            ("title", self.title_col),
            ("note", self.note_col),
            ("flag", self.flag_col),
        ):
            if col is None:
                continue
            s = _cell_str(self.value(row_1based, col))
            if s is not None:
                out[key] = s
        raw = self.value(row_1based, col_0based)
        num = _as_number(raw)
        if num is not None:
            out["v"] = num
        return out


def _as_number(raw) -> float | None:
    """Numeric cell value, or None for blanks / non-numeric text."""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None
    try:
        return float(Decimal(s))
    except (ArithmeticError, ValueError):
        return None


def _load_sheet(path: Path, sheet_name: str) -> _Sheet:
    from openpyxl import load_workbook

    if not path.exists():
        raise FileNotFoundError(f"workbook not found: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"sheet {sheet_name!r} not found in {path.name}"
                f" (have: {wb.sheetnames})"
            )
        return _Sheet(sheet_name, list(wb[sheet_name].iter_rows(values_only=True)))
    finally:
        wb.close()


def build_preview(sheet: _Sheet, *, fact_id: str, cells: str, total: float) -> dict:
    """Build one citation's preview payload against an already-loaded sheet."""
    refs = [c.strip() for c in cells.split(",") if c.strip()]
    if not refs:
        raise ValueError(f"{fact_id}: citation has no cell refs")

    cited_rows: list[int] = []
    col_letters: str | None = None
    for ref in refs:
        letters, row = parse_cell_ref(ref)
        if col_letters is None:
            col_letters = letters
        elif letters != col_letters:
            # Every loader writes one amount column per (row grain, amount
            # type); a mixed-column citation would make "the cited column"
            # meaningless, so refuse rather than pick one.
            raise ValueError(
                f"{fact_id}: cells span multiple columns ({cells}) —"
                " the preview assumes one amount column"
            )
        if row <= sheet.header_row or row > len(sheet.rows):
            raise ValueError(
                f"{fact_id}: cell {ref} names row {row}, outside the sheet body"
                f" (header row {sheet.header_row}, {len(sheet.rows)} rows)"
            )
        cited_rows.append(row)

    assert col_letters is not None
    col_0 = column_letters_to_index(col_letters) - 1

    # Decimal-exact agreement with the citation — the same comparison
    # verify_phase5b1._verify_workbook makes against the same file.
    got = Decimal(0)
    for row in cited_rows:
        raw = sheet.value(row, col_0)
        if raw is None or str(raw).strip() == "":
            raise ValueError(
                f"{fact_id}: cited cell {col_letters}{row} is blank"
            )
        try:
            got += Decimal(str(raw))
        except ArithmeticError as e:
            raise ValueError(
                f"{fact_id}: cited cell {col_letters}{row} value {raw!r}"
                f" is not numeric: {e}"
            ) from e
    want = Decimal(str(total))
    if got != want:
        raise ValueError(
            f"{fact_id}: cell sum mismatch — cells={refs} sum={got}"
            f" citation amount_thousands={want}"
        )

    lo = max(sheet.header_row + 1, min(cited_rows) - CONTEXT_ROWS)
    hi = min(len(sheet.rows), max(cited_rows) + CONTEXT_ROWS)
    cited_set = set(cited_rows)

    rows: list[dict] = []
    for r in range(lo, hi + 1):
        entry = sheet.label_row(r, col_0)
        if r in cited_set:
            entry["cited"] = True
        rows.append(entry)

    return {
        "sheet": sheet.name,
        "col": col_letters,
        "col_header": sheet.headers.get(col_0),
        "total": float(total),
        "rows": rows,
    }


def build_workbook_previews(
    *,
    workbook_dir,
    citations: dict[str, dict],
    missing_workbooks: set[str] | None = None,
) -> dict[str, dict]:
    """Preview payloads for every workbook citation in `citations`.

    `citations` is the exporter's citations_dict (fact_id -> row dict).
    Workbook files are read from `workbook_dir` as {sha256}.xlsx — the
    sha-named copies export_site already wrote.

    `missing_workbooks`: pass a set to TOLERATE .xlsx files that were never
    exported — their shas are collected here and their citations ship without
    a preview (see the module docstring: the missing document has its own
    gate, and gate 4 leg 3a catches the previewless citation). Pass nothing
    and a missing file raises FileNotFoundError.
    """
    workbook_dir = Path(workbook_dir)
    sheets: dict[tuple[str, str], _Sheet] = {}
    previews: dict[str, dict] = {}

    for fact_id in sorted(citations):
        cit = citations[fact_id]
        if cit.get("kind") != "workbook":
            continue
        if missing_workbooks is not None and cit.get("sha256") in missing_workbooks:
            continue
        sha = cit.get("sha256")
        sheet_name = cit.get("sheet")
        cells = cit.get("cells")
        total = cit.get("amount_thousands")
        if not sha or not sheet_name or not cells or total is None:
            raise ValueError(
                f"{fact_id}: workbook citation missing sha256/sheet/cells/amount"
                f" (sha={sha!r} sheet={sheet_name!r} cells={cells!r}"
                f" amount={total!r})"
            )

        key = (sha, sheet_name)
        sheet = sheets.get(key)
        if sheet is None:
            path = workbook_dir / f"{sha}.xlsx"
            if missing_workbooks is not None and not path.exists():
                missing_workbooks.add(sha)
                continue
            sheet = _load_sheet(path, sheet_name)
            sheets[key] = sheet

        payload = build_preview(
            sheet, fact_id=fact_id, cells=cells, total=float(total)
        )
        payload["units"] = cit.get("units")
        previews[fact_id] = payload

    return previews


def shard_workbook_previews(previews: dict[str, dict], out_dir) -> int:
    """Write json/workbook-cells/{fact_id[:2]}.json. Returns files written.

    Same sharding as cite-shards (256 buckets keyed by the fact_id's first
    two hex chars) so the drawer's fetch pattern is the one already proven
    by lib/cite-shards.ts.
    """
    import json

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    shards: dict[str, dict] = {}
    for fact_id, payload in previews.items():
        shards.setdefault(fact_id[:2], {})[fact_id] = payload

    for prefix, rows in shards.items():
        (out_dir / f"{prefix}.json").write_text(
            json.dumps(rows, sort_keys=True, separators=(",", ":")) + "\n",
            encoding="utf-8",
        )
    return len(shards)
