import re
from decimal import Decimal
from pathlib import Path

import psycopg
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ID_HEADERS = {
    "Account": "account",
    "Account Title": "account_title",
    "Organization": "organization",
    "Budget Activity": "budget_activity",
    "Budget Activity Title": "budget_activity_title",
    "Line Number": "line_number",
    "PE/BLI": "pe_bli",
    "Program Element/Budget Line Item (BLI) Title": "title",
}
REQUIRED = {"Account", "Organization", "PE/BLI"}


def _slug(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", header.lower()).strip("_")


def _str(v) -> str | None:
    return None if v is None else str(v)


def _find_header_row(ws) -> tuple[int, dict[int, str]]:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        cells = {j: str(v).strip() for j, v in enumerate(row) if v is not None}
        if REQUIRED <= set(cells.values()):
            return i, cells
    raise ValueError(f"No header row with {REQUIRED} found in first 20 rows")


def load_rollup(
    dsn: str, xlsx_path: Path, *, exhibit: str, fiscal_year: int,
    source_document_id: int | None = None,
) -> int:
    """Melt an R-1/P-1 display workbook into budget_lines rows. Returns upsert executions fired (inserts + updates)."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = wb[f"Exhibit {exhibit}"] if f"Exhibit {exhibit}" in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row, headers = _find_header_row(sheet)
    fy_cols = {j: _slug(h) for j, h in headers.items() if h.upper().startswith("FY ")}
    id_cols = {j: ID_HEADERS[h] for j, h in headers.items() if h in ID_HEADERS}
    upserted = 0
    with psycopg.connect(dsn) as con:
        for row_idx, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            ids = {name: row[j] for j, name in id_cols.items() if j < len(row)}
            if not ids.get("pe_bli"):
                continue
            for j, amount_type in fy_cols.items():
                if j >= len(row) or row[j] is None or str(row[j]).strip() == "":
                    continue
                try:
                    amount = Decimal(str(row[j]))
                except ArithmeticError:
                    continue
                cell = f"{get_column_letter(j + 1)}{row_idx}"
                con.execute(
                    """
                    insert into budget_lines
                      (exhibit, fiscal_year, account, account_title, organization,
                       budget_activity, budget_activity_title, line_number, pe_bli,
                       title, amount_type, amount_thousands, source_document_id,
                       source_sheet, source_cells)
                    values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                    on conflict (exhibit, fiscal_year, account, organization, budget_activity, pe_bli, amount_type)
                    do update set amount_thousands = excluded.amount_thousands,
                                  title = excluded.title,
                                  source_document_id = excluded.source_document_id,
                                  source_sheet = excluded.source_sheet,
                                  source_cells = excluded.source_cells
                    """,
                    (
                        exhibit, fiscal_year, str(ids.get("account")), ids.get("account_title"),
                        str(ids.get("organization")), _str(ids.get("budget_activity")),
                        ids.get("budget_activity_title"), _str(ids.get("line_number")),
                        str(ids.get("pe_bli")), ids.get("title"), amount_type, amount,
                        source_document_id, sheet.title, [cell],
                    ),
                )
                upserted += 1
    return upserted
