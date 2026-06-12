"""Tests for verify_phase5b1 gates.

Path-only gates — no Postgres access. Fixture site dirs built in tmp_path.
The one real-PDF jbook_pdf citation test uses tests/fixtures/jbooks/darpa_p24_25.pdf
and obtains the stored bbox by calling find_fact_page live.

Gate 1: citation_gate5b1 — stratified sample, 100% re-derivation required.
Gate 2: integrity_gate5b1 — per-kind set checks + manifest reconciliation.
Gate 3: coverage_report5b1 — non-gating coverage stats.
"""
from __future__ import annotations

import hashlib
import json
import shutil
from decimal import Decimal
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook

from govbudget.export_site import (
    fact_id_jbook,
    fact_id_lda,
    fact_id_workbook,
)
from govbudget.verify_phase5b1 import (
    citation_gate5b1,
    coverage_report5b1,
    integrity_gate5b1,
)

FIXTURE_PDF = Path(__file__).resolve().parent / "fixtures" / "jbooks" / "darpa_p24_25.pdf"

# ---------------------------------------------------------------------------
# Helpers — build minimal fixture site dirs
# ---------------------------------------------------------------------------


def _write_parquet(path: Path, col_defs: str, rows: list[tuple]) -> None:
    """Write a typed parquet via duckdb create-table + executemany."""
    path.parent.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect()
    con.execute(f"create table _t ({col_defs})")
    if rows:
        placeholders = ", ".join("?" for _ in rows[0])
        con.executemany(f"insert into _t values ({placeholders})", rows)
    path_str = str(path).replace("'", "''")
    con.execute(f"copy _t to '{path_str}' (format parquet, compression zstd)")
    con.close()


def _make_workbook(path: Path, sheet_name: str = "Exhibit R-1") -> None:
    """Create a minimal XLSX with one data row."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    # Header row 1
    ws.append(["Account", "Account Title", "Organization", "Budget Activity",
               "Budget Activity Title", "Line Number", "PE/BLI", "Title",
               "Include In TOA", "FY 2024 Actuals"])
    # Data row 2 (J2 = 280494)
    ws.append(["0400", "RDT&E Defense-Wide", "DARPA", "01", "Basic Research",
               "2", "0601101E", "DEFENSE RESEARCH SCIENCES", "Y", 280494])
    wb.save(path)


def _make_site_with_jbook_pdf(
    site_dir: Path,
    *,
    pdf_path: Path = FIXTURE_PDF,
) -> tuple[str, int, dict]:
    """Build a minimal site dir with a real jbook_pdf citation.

    Returns (sha256, page_number, bbox_dict) from the live find_fact_page call.
    """
    from govbudget.jbooks.provenance_pages import find_fact_page

    sha = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    amount = Decimal("280.494")
    hit = find_fact_page(pdf_path, pe_bli="0601101E", amount=amount)
    assert hit["resolution"] in ("unique", "ambiguous_first"), \
        f"fixture resolution unexpected: {hit['resolution']}"

    page_n = hit["page_number"]
    x0, top_pt = hit["x0"], hit["top_pt"]
    amount_text = hit["amount_text"]

    # sha-named PDF copy
    pdfs_dir = site_dir / "pdfs"
    pdfs_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(pdf_path, pdfs_dir / f"{sha}.pdf")

    fid = fact_id_jbook(sha, "0601101E", None, "PriorYear", "280.494")

    # jbook_details.parquet
    _write_parquet(
        site_dir / "data" / "jbook_details.parquet",
        "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
        " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
        " org varchar, exhibit_family varchar, fiscal_year integer,"
        " document_sha256 varchar, resolution varchar",
        [(fid, "0601101E", None, "Defense Research Sciences", "PriorYear",
          280.494, "USD millions", "ProgramElement[0]",
          "DARPA", "rdte", 2026, sha, hit["resolution"])],
    )

    # citations.parquet with jbook_pdf row
    _write_parquet(
        site_dir / "citations" / "citations.parquet",
        "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
        " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
        " page_width double, page_height double, resolution varchar,"
        " sheet varchar, cells varchar, amount_thousands double,"
        " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
        " xml_path varchar, retrieved_at varchar",
        [(fid, "jbook_pdf", "USD millions", amount_text, page_n,
          float(hit["x0"]), float(hit["x1"]),
          float(hit["top_pt"]), float(hit["bottom_pt"]),
          float(hit["page_width"]), float(hit["page_height"]),
          hit["resolution"],
          None, None, None,
          sha,
          f"https://cdn.example/pdfs/{sha}.pdf#page={page_n}",
          f"https://example.mil/darpa.pdf#page={page_n}",
          None, None)],
    )

    return sha, page_n, hit


def _make_site_with_workbook(site_dir: Path) -> tuple[str, str]:
    """Build site dir with a workbook citation.

    Returns (sha256, fact_id).
    """
    # Create XLSX fixture
    wb_src = site_dir / "_src_r1.xlsx"
    _make_workbook(wb_src)
    sha = hashlib.sha256(wb_src.read_bytes()).hexdigest()

    wb_dir = site_dir / "workbooks"
    wb_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(wb_src, wb_dir / f"{sha}.xlsx")

    fid = fact_id_workbook(sha, "R-1", 2026, "0400", "DARPA", "01", "0601101E", "fy_2024_actuals")

    # budget_lines.parquet
    _write_parquet(
        site_dir / "data" / "budget_lines.parquet",
        "fact_id varchar, exhibit varchar, fiscal_year integer, account varchar,"
        " account_title varchar, organization varchar, budget_activity varchar,"
        " budget_activity_title varchar, pe_bli varchar, title varchar,"
        " amount_type varchar, amount_thousands double, units varchar,"
        " document_sha256 varchar, source_sheet varchar, source_cells varchar",
        [(fid, "R-1", 2026, "0400", "RDT&E Defense-Wide", "DARPA", "01",
          "Basic Research", "0601101E", "DEFENSE RESEARCH SCIENCES",
          "fy_2024_actuals", 280494.0, "USD thousands",
          sha, "Exhibit R-1", "J2")],
    )

    # citations.parquet with workbook row
    _write_parquet(
        site_dir / "citations" / "citations.parquet",
        "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
        " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
        " page_width double, page_height double, resolution varchar,"
        " sheet varchar, cells varchar, amount_thousands double,"
        " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
        " xml_path varchar, retrieved_at varchar",
        [(fid, "workbook", "USD thousands", None, None,
          None, None, None, None, None, None, None,
          "Exhibit R-1", "J2", 280494.0,
          sha, None, "https://example.mil/r1.xlsx", None, None)],
    )

    return sha, fid


def _make_site_with_lda(site_dir: Path) -> tuple[str, str]:
    """Build site dir with an lda_filing citation.

    Returns (filing_uuid, fact_id).
    """
    filing_uuid = "a1b2c3d4-e5f6-7890-abcd-ef1234567890"
    pe_bli = "0601101E"
    matched_term = "darpa"
    fid = fact_id_lda(filing_uuid, pe_bli, matched_term)
    filing_url = f"https://lda.senate.gov/filings/{filing_uuid}/"

    # fct_program_lobbying.parquet  (needed for integrity gate re-derive check)
    _write_parquet(
        site_dir / "data" / "fct_program_lobbying.parquet",
        "filing_uuid varchar, pe_bli varchar, program_title varchar,"
        " matched_term varchar, description_snippet varchar,"
        " filing_url varchar, client_name varchar,"
        " family_key varchar, filing_year varchar",
        [(filing_uuid, pe_bli, "Defense Research Sciences", matched_term,
          "mentioned DARPA", filing_url, "Lockheed Martin", "lockheed", "2025")],
    )

    _write_parquet(
        site_dir / "citations" / "citations.parquet",
        "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
        " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
        " page_width double, page_height double, resolution varchar,"
        " sheet varchar, cells varchar, amount_thousands double,"
        " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
        " xml_path varchar, retrieved_at varchar",
        [(fid, "lda_filing", None, None, None,
          None, None, None, None, None, None, None,
          None, None, None,
          None, None, filing_url, None, None)],
    )

    return filing_uuid, fid


def _write_manifest(site_dir: Path, **kwargs) -> None:
    """Write a minimal manifest.json."""
    defaults = {
        "built_at": "2026-06-12T00:00:00+00:00",
        "datasets": {},
        "citations": {},
        "skipped_unresolved": 0,
        "skipped_zero_amount": 0,
        "uncited_datasets": [],
        "pdf_base_url": "/pdfs",
        "schema_version": 1,
    }
    defaults.update(kwargs)
    (site_dir / "manifest.json").write_text(json.dumps(defaults, indent=2))


# ---------------------------------------------------------------------------
# Gate 1: citation_gate5b1 — happy paths
# ---------------------------------------------------------------------------


class TestCitationGate:
    def test_jbook_pdf_happy_path(self, tmp_path):
        """Real PDF fixture — bbox within 2 pt of stored values."""
        site = tmp_path / "site"
        sha, page_n, hit = _make_site_with_jbook_pdf(site)
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is True, f"failures: {result.get('failures')}"
        assert result["sampled"] >= 1
        assert result["passed"] == result["sampled"]
        assert result["failures"] == []

    def test_workbook_happy_path(self, tmp_path):
        """Workbook cell sum must match stored amount_thousands."""
        site = tmp_path / "site"
        _make_site_with_workbook(site)
        _write_manifest(site, datasets={"budget_lines": 1}, citations={"workbook": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is True, f"failures: {result.get('failures')}"
        assert result["sampled"] >= 1
        assert result["passed"] == result["sampled"]

    def test_lda_filing_happy_path(self, tmp_path):
        """LDA shape check — url starts with lda.senate.gov and contains uuid."""
        site = tmp_path / "site"
        filing_uuid, fid = _make_site_with_lda(site)
        _write_manifest(site, citations={"lda_filing": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is True, f"failures: {result.get('failures')}"

    def test_empty_citations_fails(self, tmp_path):
        """Empty citations.parquet → ok=False."""
        site = tmp_path / "site"
        (site / "citations").mkdir(parents=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
            " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
            " page_width double, page_height double, resolution varchar,"
            " sheet varchar, cells varchar, amount_thousands double,"
            " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
            " xml_path varchar, retrieved_at varchar",
            [],
        )
        _write_manifest(site)

        result = citation_gate5b1(site)

        assert result["ok"] is False
        assert result["sampled"] == 0

    def test_missing_site_dir_fails(self, tmp_path):
        result = citation_gate5b1(tmp_path / "nonexistent")
        assert result["ok"] is False
        assert "reason" in result

    def test_tampered_amount_text_fails(self, tmp_path):
        """Storing wrong amount_text → PDF word search fails → gate FAIL."""
        site = tmp_path / "site"
        sha, page_n, hit = _make_site_with_jbook_pdf(site)
        fid = fact_id_jbook(sha, "0601101E", None, "PriorYear", "280.494")

        # Re-write citations with a wrong amount_text
        (site / "citations").mkdir(parents=True, exist_ok=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
            " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
            " page_width double, page_height double, resolution varchar,"
            " sheet varchar, cells varchar, amount_thousands double,"
            " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
            " xml_path varchar, retrieved_at varchar",
            [(fid, "jbook_pdf", "USD millions", "999999.000", page_n,
              float(hit["x0"]), float(hit["x1"]),
              float(hit["top_pt"]), float(hit["bottom_pt"]),
              float(hit["page_width"]), float(hit["page_height"]),
              hit["resolution"],
              None, None, None,
              sha,
              f"https://cdn.example/pdfs/{sha}.pdf#page={page_n}",
              f"https://example.mil/darpa.pdf#page={page_n}",
              None, None)],
        )
        _write_manifest(site, citations={"jbook_pdf": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is False
        assert len(result["failures"]) >= 1
        assert any(fid == f[0] for f in result["failures"]), \
            f"expected {fid} in failures, got: {result['failures']}"

    def test_tampered_workbook_amount_fails(self, tmp_path):
        """Stored amount_thousands differs from cell value → gate FAIL."""
        site = tmp_path / "site"
        sha, fid = _make_site_with_workbook(site)

        # Re-write citations with wrong amount_thousands (does NOT match cell J2=280494)
        (site / "citations").mkdir(parents=True, exist_ok=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
            " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
            " page_width double, page_height double, resolution varchar,"
            " sheet varchar, cells varchar, amount_thousands double,"
            " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
            " xml_path varchar, retrieved_at varchar",
            [(fid, "workbook", "USD thousands", None, None,
              None, None, None, None, None, None, None,
              "Exhibit R-1", "J2", 999999.0,   # WRONG amount
              sha, None, "https://example.mil/r1.xlsx", None, None)],
        )
        _write_manifest(site, citations={"workbook": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is False
        assert len(result["failures"]) >= 1

    def test_lda_bad_url_fails(self, tmp_path):
        """LDA citation with non-lda.senate.gov URL → gate FAIL."""
        site = tmp_path / "site"
        fid = fact_id_lda("uuid-bad", "0601101E", "darpa")
        (site / "citations").mkdir(parents=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
            " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
            " page_width double, page_height double, resolution varchar,"
            " sheet varchar, cells varchar, amount_thousands double,"
            " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
            " xml_path varchar, retrieved_at varchar",
            [(fid, "lda_filing", None, None, None,
              None, None, None, None, None, None, None,
              None, None, None,
              None, None, "https://not-lda.example.com/f/1", None, None)],
        )
        _write_manifest(site, citations={"lda_filing": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is False
        assert len(result["failures"]) >= 1

    def test_lda_url_without_uuid_fails(self, tmp_path):
        """LDA citation whose official_url has no UUID in the path → gate FAIL."""
        site = tmp_path / "site"
        fid = fact_id_lda("no-uuid-here", "0601101E", "darpa")
        (site / "citations").mkdir(parents=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
            " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
            " page_width double, page_height double, resolution varchar,"
            " sheet varchar, cells varchar, amount_thousands double,"
            " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
            " xml_path varchar, retrieved_at varchar",
            [(fid, "lda_filing", None, None, None,
              None, None, None, None, None, None, None,
              None, None, None,
              None, None, "https://lda.senate.gov/filings/no-uuid-here/", None, None)],
        )
        _write_manifest(site, citations={"lda_filing": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is False, \
            "expected FAIL when official_url has no filing UUID"
        assert len(result["failures"]) >= 1
        # The failure message must identify the fact_id
        assert any(fid in str(f) for f in result["failures"]), \
            f"expected {fid} in failures, got: {result['failures']}"


# ---------------------------------------------------------------------------
# Stratified sampling guarantee
# ---------------------------------------------------------------------------


class TestStratifiedSampling:
    def test_stratification_guarantees_min_per_kind(self, tmp_path):
        """100 jbook rows + 2 workbook + 2 lda, sample_size=50.

        Every kind must appear in sample; workbook and lda must each get
        ALL their rows (2 each); jbook gets the remainder (46).

        The test detects the greedy bug by monkey-patching _verify_* to record
        which fact_ids were actually sampled, then checking per-kind counts.
        """
        site = tmp_path / "site"

        # Build a single valid workbook fixture (for the workbook citations)
        wb_src = site / "_r1.xlsx"
        _make_workbook(wb_src)
        sha_wb = hashlib.sha256(wb_src.read_bytes()).hexdigest()
        (site / "workbooks").mkdir(parents=True, exist_ok=True)
        shutil.copyfile(wb_src, site / "workbooks" / f"{sha_wb}.xlsx")

        # Build a single valid PDF fixture (for the jbook citations)
        sha_pdf = hashlib.sha256(FIXTURE_PDF.read_bytes()).hexdigest()
        (site / "pdfs").mkdir(parents=True, exist_ok=True)
        shutil.copyfile(FIXTURE_PDF, site / "pdfs" / f"{sha_pdf}.pdf")

        from govbudget.jbooks.provenance_pages import find_fact_page
        hit = find_fact_page(FIXTURE_PDF, pe_bli="0601101E", amount=Decimal("280.494"))

        # 100 jbook_pdf rows (same sha/pe_bli/page — same physical citation; fact_id
        # differs by a unique suffix baked into fact_id_jbook via scenario field)
        jbook_rows = []
        for i in range(100):
            fid = fact_id_jbook(sha_pdf, "0601101E", None, f"Scenario{i:04d}", "280.494")
            jbook_rows.append((
                fid, "jbook_pdf", "USD millions", hit["amount_text"],
                hit["page_number"],
                float(hit["x0"]), float(hit["x1"]),
                float(hit["top_pt"]), float(hit["bottom_pt"]),
                float(hit["page_width"]), float(hit["page_height"]),
                hit["resolution"],
                None, None, None,
                sha_pdf,
                f"https://cdn.example/pdfs/{sha_pdf}.pdf#page={hit['page_number']}",
                f"https://example.mil/darpa.pdf#page={hit['page_number']}",
                None, None,
            ))

        # 2 workbook rows
        wb_rows = []
        for i in range(2):
            fid = fact_id_workbook(sha_wb, "R-1", 2026, "0400", f"ORG{i}", "01", "0601101E", "fy_2024_actuals")
            wb_rows.append((
                fid, "workbook", "USD thousands", None, None,
                None, None, None, None, None, None, None,
                "Exhibit R-1", "J2", 280494.0,
                sha_wb, None, "https://example.mil/r1.xlsx", None, None,
            ))

        # 2 lda_filing rows (use a real UUID so the UUID check passes)
        lda_rows = []
        real_uuid = "a1b2c3d4-e5f6-7890-abcd-ef1234567890"
        for i in range(2):
            fid = fact_id_lda(f"uuid-lda-{i:03d}", "0601101E", "darpa")
            lda_rows.append((
                fid, "lda_filing", None, None, None,
                None, None, None, None, None, None, None,
                None, None, None,
                None, None, f"https://lda.senate.gov/filings/{real_uuid}/", None, None,
            ))

        all_rows = jbook_rows + wb_rows + lda_rows

        col_defs = (
            "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
            " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
            " page_width double, page_height double, resolution varchar,"
            " sheet varchar, cells varchar, amount_thousands double,"
            " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
            " xml_path varchar, retrieved_at varchar"
        )
        _write_parquet(site / "citations" / "citations.parquet", col_defs, all_rows)
        _write_manifest(site)

        # Build a fact_id → kind lookup for asserting per-kind coverage
        wb_fact_ids = {r[0] for r in wb_rows}
        lda_fact_ids = {r[0] for r in lda_rows}
        jbook_fact_ids = {r[0] for r in jbook_rows}

        # Monkey-patch the _verify_* functions to record sampled fact_ids
        import govbudget.verify_phase5b1 as _mod
        sampled_fact_ids: list[str] = []
        _orig_jbook = _mod._verify_jbook_pdf
        _orig_wb = _mod._verify_workbook
        _orig_lda = _mod._verify_lda

        def _record_jbook(site_dir, row, idx):
            sampled_fact_ids.append(row[idx["fact_id"]])
            return _orig_jbook(site_dir, row, idx)

        def _record_wb(site_dir, row, idx):
            sampled_fact_ids.append(row[idx["fact_id"]])
            return _orig_wb(site_dir, row, idx)

        def _record_lda(row, idx):
            sampled_fact_ids.append(row[idx["fact_id"]])
            return _orig_lda(row, idx)

        _mod._verify_jbook_pdf = _record_jbook
        _mod._verify_workbook = _record_wb
        _mod._verify_lda = _record_lda
        try:
            result = citation_gate5b1(site, sample_size=50)
        finally:
            _mod._verify_jbook_pdf = _orig_jbook
            _mod._verify_workbook = _orig_wb
            _mod._verify_lda = _orig_lda

        assert result["ok"] is True, f"gate failures: {result.get('failures')}"
        assert result["sampled"] == 50

        sampled_set = set(sampled_fact_ids)
        n_jbook_sampled = len(sampled_set & jbook_fact_ids)
        n_wb_sampled = len(sampled_set & wb_fact_ids)
        n_lda_sampled = len(sampled_set & lda_fact_ids)

        # workbook and lda have only 2 rows each → must both be fully sampled
        assert n_wb_sampled == 2, \
            f"workbook under-sampled: got {n_wb_sampled}/2 (greedy bug?)"
        assert n_lda_sampled == 2, \
            f"lda under-sampled: got {n_lda_sampled}/2 (greedy bug?)"
        assert n_jbook_sampled == 46, \
            f"jbook count wrong: got {n_jbook_sampled}, expected 46 (= 50 - 2 - 2)"


# ---------------------------------------------------------------------------
# Gate 2: integrity_gate5b1
# ---------------------------------------------------------------------------


class TestIntegrityGate:
    def _make_full_site(self, site_dir: Path) -> dict:
        """Build a minimal but complete site with jbook, workbook, and lda citations."""
        # ---- jbook ----
        sha_pdf = hashlib.sha256(FIXTURE_PDF.read_bytes()).hexdigest()
        from govbudget.jbooks.provenance_pages import find_fact_page
        hit = find_fact_page(FIXTURE_PDF, pe_bli="0601101E", amount=Decimal("280.494"))
        page_n = hit["page_number"]
        fid_jbook = fact_id_jbook(sha_pdf, "0601101E", None, "PriorYear", "280.494")

        pdfs_dir = site_dir / "pdfs"
        pdfs_dir.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(FIXTURE_PDF, pdfs_dir / f"{sha_pdf}.pdf")

        # ---- workbook ----
        wb_src = site_dir / "_r1.xlsx"
        _make_workbook(wb_src)
        sha_wb = hashlib.sha256(wb_src.read_bytes()).hexdigest()
        wb_dir = site_dir / "workbooks"
        wb_dir.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(wb_src, wb_dir / f"{sha_wb}.xlsx")
        fid_wb = fact_id_workbook(sha_wb, "R-1", 2026, "0400", "DARPA", "01", "0601101E", "fy_2024_actuals")

        # ---- lda ----
        fid_lda = fact_id_lda("uuid-001", "0601101E", "darpa")

        # ---- data parquets ----
        _write_parquet(
            site_dir / "data" / "jbook_details.parquet",
            "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
            " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
            " org varchar, exhibit_family varchar, fiscal_year integer,"
            " document_sha256 varchar, resolution varchar",
            [(fid_jbook, "0601101E", None, "Defense Research Sciences", "PriorYear",
              280.494, "USD millions", "ProgramElement[0]",
              "DARPA", "rdte", 2026, sha_pdf, hit["resolution"])],
        )
        _write_parquet(
            site_dir / "data" / "budget_lines.parquet",
            "fact_id varchar, exhibit varchar, fiscal_year integer, account varchar,"
            " account_title varchar, organization varchar, budget_activity varchar,"
            " budget_activity_title varchar, pe_bli varchar, title varchar,"
            " amount_type varchar, amount_thousands double, units varchar,"
            " document_sha256 varchar, source_sheet varchar, source_cells varchar",
            [(fid_wb, "R-1", 2026, "0400", "RDT&E Defense-Wide", "DARPA", "01",
              "Basic Research", "0601101E", "DEFENSE RESEARCH SCIENCES",
              "fy_2024_actuals", 280494.0, "USD thousands",
              sha_wb, "Exhibit R-1", "J2")],
        )
        _write_parquet(
            site_dir / "data" / "fct_program_lobbying.parquet",
            "filing_uuid varchar, pe_bli varchar, program_title varchar,"
            " matched_term varchar, description_snippet varchar,"
            " filing_url varchar, client_name varchar,"
            " family_key varchar, filing_year varchar",
            [("uuid-001", "0601101E", "Defense Research Sciences", "darpa",
              "mentioned DARPA", "https://lda.senate.gov/filings/uuid-001/",
              "Lockheed Martin", "lockheed", "2025")],
        )

        # ---- citations.parquet ----
        cit_rows = [
            (fid_jbook, "jbook_pdf", "USD millions", hit["amount_text"],
             page_n, float(hit["x0"]), float(hit["x1"]),
             float(hit["top_pt"]), float(hit["bottom_pt"]),
             float(hit["page_width"]), float(hit["page_height"]),
             hit["resolution"],
             None, None, None,
             sha_pdf,
             f"https://cdn.example/pdfs/{sha_pdf}.pdf#page={page_n}",
             f"https://example.mil/darpa.pdf#page={page_n}",
             None, None),
            (fid_wb, "workbook", "USD thousands", None, None,
             None, None, None, None, None, None, None,
             "Exhibit R-1", "J2", 280494.0,
             sha_wb, None, "https://example.mil/r1.xlsx", None, None),
            (fid_lda, "lda_filing", None, None, None,
             None, None, None, None, None, None, None,
             None, None, None,
             None, None, "https://lda.senate.gov/filings/uuid-001/", None, None),
        ]
        _write_parquet(
            site_dir / "citations" / "citations.parquet",
            "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
            " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
            " page_width double, page_height double, resolution varchar,"
            " sheet varchar, cells varchar, amount_thousands double,"
            " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
            " xml_path varchar, retrieved_at varchar",
            cit_rows,
        )

        # ---- manifest.json ----
        _write_manifest(
            site_dir,
            datasets={"jbook_details": 1, "budget_lines": 1, "fct_program_lobbying": 1},
            citations={"jbook_pdf": 1, "workbook": 1, "lda_filing": 1},
            skipped_unresolved=0,
            skipped_zero_amount=0,
        )

        return {
            "fid_jbook": fid_jbook, "fid_wb": fid_wb, "fid_lda": fid_lda,
            "sha_pdf": sha_pdf, "sha_wb": sha_wb,
        }

    def test_happy_path(self, tmp_path):
        site = tmp_path / "site"
        self._make_full_site(site)
        result = integrity_gate5b1(site)
        assert result["ok"] is True, f"failures: {result}"

    def test_orphan_jbook_citation_fails(self, tmp_path):
        """A jbook_pdf citation whose fact_id is NOT in jbook_details → FAIL."""
        site = tmp_path / "site"
        ids = self._make_full_site(site)

        # Add an extra jbook_pdf citation with an orphan fact_id
        orphan_fid = "orphan0000000000"
        sha = ids["sha_pdf"]
        cit_pq = site / "citations" / "citations.parquet"
        con = duckdb.connect()
        existing = con.execute(f"select * from read_parquet('{cit_pq}')").fetchall()
        cols = [d[0] for d in con.execute(f"describe select * from read_parquet('{cit_pq}')").fetchall()]
        con.close()

        from govbudget.jbooks.provenance_pages import find_fact_page
        hit = find_fact_page(FIXTURE_PDF, pe_bli="0601101E", amount=Decimal("280.494"))
        extra_row = (orphan_fid, "jbook_pdf", "USD millions", hit["amount_text"],
                     hit["page_number"],
                     float(hit["x0"]), float(hit["x1"]),
                     float(hit["top_pt"]), float(hit["bottom_pt"]),
                     float(hit["page_width"]), float(hit["page_height"]),
                     hit["resolution"],
                     None, None, None,
                     sha,
                     f"https://cdn.example/pdfs/{sha}.pdf#page={hit['page_number']}",
                     "https://example.mil/darpa.pdf#page=1",
                     None, None)
        _write_parquet(cit_pq, ", ".join(f"{c} varchar" for c in cols), existing + [extra_row])

        result = integrity_gate5b1(site)
        assert result["ok"] is False

    def test_resolved_details_row_missing_citation_fails(self, tmp_path):
        """A jbook_details row with resolution='unique' missing from citations → FAIL."""
        site = tmp_path / "site"
        ids = self._make_full_site(site)

        # Add another details row with a new fact_id (resolution=unique) that has no citation
        details_pq = site / "data" / "jbook_details.parquet"
        missing_fid = "missing00000000"
        con = duckdb.connect()
        existing = con.execute(f"select * from read_parquet('{details_pq}')").fetchall()
        extra = (missing_fid, "0602303E", None, "ICT Research", "BudgetYearOne",
                 100.0, "USD millions", "PE[1]", "DARPA", "rdte", 2026,
                 ids["sha_pdf"], "unique")
        _write_parquet(
            details_pq,
            "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
            " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
            " org varchar, exhibit_family varchar, fiscal_year integer,"
            " document_sha256 varchar, resolution varchar",
            existing + [extra],
        )
        con.close()

        result = integrity_gate5b1(site)
        assert result["ok"] is False

    def test_manifest_rowcount_mismatch_fails(self, tmp_path):
        """manifest.json dataset count that doesn't match actual parquet → FAIL."""
        site = tmp_path / "site"
        self._make_full_site(site)

        # Overwrite manifest with wrong count for jbook_details
        _write_manifest(
            site,
            datasets={"jbook_details": 999, "budget_lines": 1, "fct_program_lobbying": 1},
            citations={"jbook_pdf": 1, "workbook": 1, "lda_filing": 1},
        )

        result = integrity_gate5b1(site)
        assert result["ok"] is False

    def test_workbook_fact_ids_mismatch_fails(self, tmp_path):
        """workbook citations contain a fact_id not in budget_lines → FAIL."""
        site = tmp_path / "site"
        ids = self._make_full_site(site)

        # Add extra workbook citation with unknown fact_id
        extra_fid = "extraworkbook00"
        cit_pq = site / "citations" / "citations.parquet"
        con = duckdb.connect()
        existing = con.execute(f"select * from read_parquet('{cit_pq}')").fetchall()
        con.close()
        extra = (extra_fid, "workbook", "USD thousands", None, None,
                 None, None, None, None, None, None, None,
                 "Sheet1", "A1", 100.0,
                 ids["sha_wb"], None, "https://example.mil/r1.xlsx", None, None)
        _write_parquet(
            cit_pq,
            "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
            " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
            " page_width double, page_height double, resolution varchar,"
            " sheet varchar, cells varchar, amount_thousands double,"
            " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
            " xml_path varchar, retrieved_at varchar",
            existing + [extra],
        )

        result = integrity_gate5b1(site)
        assert result["ok"] is False

    def test_missing_site_dir_fails(self, tmp_path):
        result = integrity_gate5b1(tmp_path / "nonexistent")
        assert result["ok"] is False

    def test_manifest_skip_count_mismatch_fails(self, tmp_path):
        """manifest skipped_unresolved is higher than parquet count → integrity FAIL."""
        site = tmp_path / "site"
        self._make_full_site(site)

        # Tamper: set skipped_unresolved to 1 higher than actual (parquet has 0 unresolved)
        man = json.loads((site / "manifest.json").read_text())
        man["skipped_unresolved"] = man.get("skipped_unresolved", 0) + 1
        (site / "manifest.json").write_text(json.dumps(man, indent=2))

        result = integrity_gate5b1(site)
        assert result["ok"] is False, \
            "expected FAIL when manifest skipped_unresolved doesn't match parquet counts"
        assert any("skip" in f.lower() or "skipped" in f.lower() or "unresolved" in f.lower()
                   for f in result["failures"]), \
            f"expected skip-count failure description, got: {result['failures']}"


# ---------------------------------------------------------------------------
# Gate 3: coverage_report5b1 (non-gating)
# ---------------------------------------------------------------------------


class TestCoverageReport:
    def test_returns_non_failing_dict(self, tmp_path):
        site = tmp_path / "site"
        sha = hashlib.sha256(FIXTURE_PDF.read_bytes()).hexdigest()
        fid = fact_id_jbook(sha, "0601101E", None, "PriorYear", "280.494")
        _write_parquet(
            site / "data" / "jbook_details.parquet",
            "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
            " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
            " org varchar, exhibit_family varchar, fiscal_year integer,"
            " document_sha256 varchar, resolution varchar",
            [(fid, "0601101E", None, "DR Sciences", "PriorYear",
              280.494, "USD millions", "PE[0]", "DARPA", "rdte", 2026,
              sha, "unique"),
             ("fid2", "0601102E", None, "Another", "BudgetYearOne",
              0.0, "USD millions", "PE[1]", "DARPA", "rdte", 2026,
              sha, "zero_amount"),
             ("fid3", "0601103E", None, "Yet Another", "BudgetYearOne",
              50.0, "USD millions", "PE[2]", "DARPA", "rdte", 2026,
              sha, "unresolved")],
        )
        _write_manifest(site, datasets={"jbook_details": 3}, uncited_datasets=["fct_budget_trajectory"])

        report = coverage_report5b1(site)

        # Non-gating: always returns a dict (never raises on valid input)
        assert isinstance(report, dict)
        assert "unique" in report or "resolved" in report or "by_resolution" in report
        # uncited_datasets echoed
        assert "uncited_datasets" in report

    def test_missing_site_dir_returns_dict(self, tmp_path):
        """Non-gating: missing dir returns a dict with ok=False or a summary."""
        report = coverage_report5b1(tmp_path / "nonexistent")
        assert isinstance(report, dict)
