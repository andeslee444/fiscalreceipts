from decimal import Decimal

import psycopg
from openpyxl import Workbook

from govbudget.jbooks.p1_loader import load_p1_rollup

P1_HEADERS = [
    "Account", "Account Title", "Organization", "Budget Activity",
    "Budget Activity Title", "Line Number", "BSA", "Budget SubActivity (BSA) Title",
    "Budget Line Item", "Budget Line Item (BLI) Title", "Cost Type",
    "Cost Type Title", "Add/Non-Add",
    "FY 2024 Actuals Quantity", "FY 2024 Actuals Amount",
    "FY 2025 Enacted Quantity", "FY 2025 Enacted Amount",
    "FY 2026 Disc Request Quantity", "FY 2026 Disc Request Amount",
    "Classification",
]


def make_p1_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit P-1"
    ws.append(["Total of Displayed Rows"])  # banner
    ws.append(P1_HEADERS)
    # same BLI, two Add cost-type rows + one Non-Add row (must be excluded)
    ws.append(["0300D", "Procurement, Defense-Wide", "CBDP", "03", "Chem/Bio Defense",
               "120", "1", "CBDP", "7001SA1000", "CB Situational Awareness",
               "A", "Weapon System Cost", "Add", 2, 100000, 1, 150000, "", 200051])
    ws.append(["0300D", "Procurement, Defense-Wide", "CBDP", "03", "Chem/Bio Defense",
               "120", "1", "CBDP", "7001SA1000", "CB Situational Awareness",
               "B", "Advance Procurement", "Add", "", 48640, "", 36841, "", 8000])
    ws.append(["0300D", "Procurement, Defense-Wide", "CBDP", "03", "Chem/Bio Defense",
               "120", "1", "CBDP", "7001SA1000", "CB Situational Awareness",
               "Z", "Non-add memo", "Non-Add", "", 999999, "", 999999, "", 999999])
    p = tmp_path / "p1_display.xlsx"
    wb.save(p)
    return p


def test_p1_loader_aggregates_to_bli(pg_dsn, tmp_path):
    n = load_p1_rollup(pg_dsn, make_p1_xlsx(tmp_path), exhibit="P-1", fiscal_year=2026)
    assert n == 3  # one BLI x three FY amount columns
    with psycopg.connect(pg_dsn) as con:
        rows = dict(con.execute(
            "select amount_type, amount_thousands from budget_lines"
            " where exhibit='P-1' and pe_bli='7001SA1000'"
        ).fetchall())
    assert rows["fy_2024_actuals"] == Decimal("148640")  # 100000 + 48640, Non-Add excluded
    assert rows["fy_2025_enacted"] == Decimal("186841")
    assert rows["fy_2026_disc_request"] == Decimal("208051")


def test_p1_loader_is_idempotent(pg_dsn, tmp_path):
    p = make_p1_xlsx(tmp_path)
    load_p1_rollup(pg_dsn, p, exhibit="P-1", fiscal_year=2026)
    load_p1_rollup(pg_dsn, p, exhibit="P-1", fiscal_year=2026)
    with psycopg.connect(pg_dsn) as con:
        total = con.execute(
            "select count(*) from budget_lines where exhibit='P-1'"
        ).fetchone()[0]
    assert total == 3


def test_load_p1_records_contributing_cells(pg_dsn, tmp_path):
    xlsx = make_p1_xlsx(tmp_path)
    load_p1_rollup(pg_dsn, xlsx, exhibit="P-1", fiscal_year=2026)
    with psycopg.connect(pg_dsn) as con:
        row = con.execute(
            "select source_sheet, source_cells from budget_lines where pe_bli = %s"
            " and amount_type = 'fy_2024_actuals'",
            ("7001SA1000",),
        ).fetchone()
    assert row[1] == ["O3", "O4"]      # BOTH contributing cells — honesty for summed facts


def test_p1_loader_skips_nan_amount(pg_dsn, tmp_path):
    """Cells whose value is the literal string 'NaN' produce Decimal('NaN') without raising.
    For P-1, this NaN contaminates the sums bucket. Must be caught by is_nan() guard
    so both the sum and the cell list are excluded together."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit P-1"
    ws.append(["banner"])
    ws.append(P1_HEADERS)
    # Row with one "NaN" amount cell and one valid amount cell
    ws.append(["0300D", "Procurement, Defense-Wide", "CBDP", "03", "Chem/Bio Defense",
               "120", "1", "CBDP", "9999NAN1", "NaN Test BLI",
               "A", "Weapon System Cost", "Add",
               2, "NaN",  # FY 2024 Actuals Amount = "NaN" → Decimal('NaN') without raise
               1, 150000,  # FY 2025 Enacted Amount = valid
               "", 0])     # FY 2026 Disc Request Amount = 0 (valid)
    p = tmp_path / "p1_nan.xlsx"
    wb.save(p)

    load_p1_rollup(pg_dsn, p, exhibit="P-1", fiscal_year=2026)
    with psycopg.connect(pg_dsn) as con:
        rows = dict(con.execute(
            "select amount_type, amount_thousands from budget_lines"
            " where exhibit='P-1' and pe_bli='9999NAN1'"
        ).fetchall())
    # NaN cell must be entirely absent (no budget_lines row for fy_2024_actuals)
    assert "fy_2024_actuals" not in rows, "NaN amount must not produce a budget_lines row"
    # Valid amounts must still load
    assert rows["fy_2025_enacted"] == Decimal("150000")
    assert rows["fy_2026_disc_request"] == Decimal("0")


def test_same_bli_different_line_numbers_sum_not_overwrite(pg_dsn, tmp_path):
    # Real-data regression: Apache-style BLI split across two Line Numbers
    # previously collided on upsert (last write wins, $138B understated).
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit P-1"
    ws.append(P1_HEADERS)
    ws.append(["2031A", "Aircraft Procurement, Army", "A", "01", "Aircraft",
               "9", "20", "Rotary", "5757A05111", "AH-64 Apache Reman",
               "A", "Weapon System Cost", "Add", 37, 649000, "", "", "", ""])
    ws.append(["2031A", "Aircraft Procurement, Army", "A", "01", "Aircraft",
               "10", "20", "Rotary", "5757A05111", "AH-64 Apache Reman",
               "A", "Advance Procurement", "Add", "", 110000, "", "", "", ""])
    p = tmp_path / "p1_split_lines.xlsx"
    wb.save(p)

    load_p1_rollup(pg_dsn, p, exhibit="P-1", fiscal_year=2026)
    with psycopg.connect(pg_dsn) as con:
        total = con.execute(
            "select sum(amount_thousands) from budget_lines"
            " where pe_bli='5757A05111' and amount_type='fy_2024_actuals'"
        ).fetchone()[0]
        n_rows = con.execute(
            "select count(*) from budget_lines"
            " where pe_bli='5757A05111' and amount_type='fy_2024_actuals'"
        ).fetchone()[0]
    assert total == Decimal("759000")  # 649000 + 110000
    assert n_rows == 1  # single aggregated control row


def test_p1_loader_handles_footnote_asterisk_headers(pg_dsn, tmp_path):
    """PB2025's P-1 labels the FY2024 column 'FY 2024 PB Request with CR
    Adjustments Amount*' — the footnote asterisk must not drop the column
    (Task 4 live-run evidence)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit P-1"
    ws.append(["Total of Displayed Rows"])
    ws.append([
        "Account", "Account Title", "Organization", "Budget Activity",
        "Budget Activity Title", "Line Number", "Budget Line Item",
        "Budget Line Item (BLI) Title", "Add/Non-Add",
        "FY 2023 Actuals Quantity", "FY 2023 Actuals Amount",
        "FY 2024 PB Request with CR Adjustments Quantity",
        "FY 2024 PB Request with CR Adjustments Amount*",
        "FY 2025 Request Quantity", "FY 2025 Request Amount",
    ])
    ws.append(["0300D", "Procurement, Defense-Wide", "CBDP", "03", "Chem/Bio",
               "120", "7001SA1000", "CB Situational Awareness", "Add",
               2, 100000, 1, 150000, 3, 200000])
    p = tmp_path / "p1_display.xlsx"
    wb.save(p)
    n = load_p1_rollup(pg_dsn, p, exhibit="P-1", fiscal_year=2025)
    assert n == 3
    with psycopg.connect(pg_dsn) as con:
        rows = dict(con.execute(
            "select amount_type, amount_thousands from budget_lines"
            " where exhibit='P-1' and pe_bli='7001SA1000'"
        ).fetchall())
    assert rows["fy_2023_actuals"] == Decimal("100000")
    assert rows["fy_2024_pb_request_with_cr_adjustments"] == Decimal("150000")
    assert rows["fy_2025_request"] == Decimal("200000")


# --- PB2017–PB2023 era header variants (live workbook evidence, Task 4) ---

ERA_P1_HEADERS = [
    "Account", "Account Title", "Organization", "Budget\nActivity",
    "Budget Activity Title", "Line\nNumber", "BSA",
    "Budget Sub Activity (BSA) Title", "Line Item", "Line Item Title",
    "Cost\nType", "Cost Type Title", "Add/\nNon-Add",
    "FY 2016\nBase Enacted\nQuantity", "FY 2016\nBase Enacted\nAmount",
    "FY 2016\nTotal Enacted\nQuantity", "FY 2016\nTotal Enacted\nAmount",
    "FY 2017\nTotal\nQuantity", "FY 2017\nTotal\nAmount", "Classification",
]


def test_p1_loader_era_headers_line_item_and_wrapped_add(pg_dsn, tmp_path):
    """PB2017–PB2023 P-1 workbooks say 'Line Item' (not 'Budget Line Item'),
    wrap 'Add/Non-Add' and the FY '... Amount' suffixes across lines."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit P-1"
    ws.append(["Total of Displayed Rows"])
    ws.append(ERA_P1_HEADERS)
    ws.append(["0300D", "Procurement, Defense-Wide", "DTRA", "01", "Major Equipment",
               "14 ", "1", "Major Equipment, DTRA", "23", "Vehicles",
               "A", "Weapon System Cost", "Add", "", 12000, "", 13000, "", 14000])
    ws.append(["0300D", "Procurement, Defense-Wide", "DTRA", "01", "Major Equipment",
               "14 ", "1", "Major Equipment, DTRA", "23", "Vehicles",
               "Z", "Memo", "Non-Add", "", 999999, "", 999999, "", 999999])
    p = tmp_path / "p1_display.xlsx"
    wb.save(p)
    n = load_p1_rollup(pg_dsn, p, exhibit="P-1", fiscal_year=2017)
    assert n == 3
    # era workbooks key budget_lines on the P-1 line number (the only
    # identifier the era P-40 XML shares — its P1LineNumber), NOT the
    # 'Line Item' display code. The line number is namespaced to
    # '{account}-{org}-L{line}' (Finding D: bare line numbers collide with
    # modern BLI codes and conflate programs across orgs).
    with psycopg.connect(pg_dsn) as con:
        rows = dict(con.execute(
            "select amount_type, amount_thousands from budget_lines"
            " where exhibit='P-1' and pe_bli='0300D-DTRA-L14'"
        ).fetchall())
        li_rows = con.execute(
            "select count(*) from budget_lines where pe_bli in ('23', '14')"
        ).fetchone()[0]
    assert rows == {
        "fy_2016_base_enacted": Decimal("12000"),
        "fy_2016_total_enacted": Decimal("13000"),
        "fy_2017_total": Decimal("14000"),
    }
    assert li_rows == 0


def test_p1r_loader_era_without_add_non_add_column(pg_dsn, tmp_path):
    """PB2017–PB2023 P-1R workbooks have no Add/Non-Add column: every row
    loads (there is nothing to filter on)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit P-1R"
    ws.append(["Total of Displayed Rows"])
    ws.append(["Account", "Account Title", "Organization", "Budget\nActivity",
               "Budget Activity Title", "BSA", "Budget Sub Activity (BSA) Title",
               "Line Item", "Line Item Title", "Cost\nType", "Cost Type Title",
               "FY 2018\nQuantity", "FY 2018\nAmount", "Classification"])
    ws.append(["0300D", "Procurement, Defense-Wide", "DEFW", "01", "Major Equipment",
               "1", "Major Equipment", "23", "Vehicles", "A", "Weapon System Cost",
               "", 21000, "U"])
    p = tmp_path / "p1r_display.xlsx"
    wb.save(p)
    n = load_p1_rollup(pg_dsn, p, exhibit="P-1R", fiscal_year=2017)
    assert n == 1
    with psycopg.connect(pg_dsn) as con:
        row = con.execute(
            "select amount_type, amount_thousands from budget_lines"
            " where exhibit='P-1R' and pe_bli='23'"
        ).fetchone()
    assert row == ("fy_2018", Decimal("21000"))
