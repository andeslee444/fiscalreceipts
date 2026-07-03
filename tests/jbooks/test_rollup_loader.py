from decimal import Decimal

import psycopg
from openpyxl import Workbook

from govbudget.jbooks.rollup_loader import load_rollup

HEADERS = [
    "Account", "Account Title", "Organization", "Budget Activity",
    "Budget Activity Title", "Line Number", "PE/BLI",
    "Program Element/Budget Line Item (BLI) Title", "Include In TOA",
    "FY 2024 Actuals", "FY 2025 Enacted", "FY 2025 Total", "FY 2026 Disc Request",
]
ROW1 = ["0400", "RDT&E Defense-Wide", "DARPA", "01", "Basic Research", "2",
        "0601101E", "DEFENSE RESEARCH SCIENCES", "Y",
        280494, 293145, 293145, 0]
ROW2 = ["0400", "RDT&E Defense-Wide", "DARPA", "02", "Applied Research", "14",
        "0602303E", "INFORMATION & COMMUNICATIONS TECHNOLOGY", "Y",
        413000, 400000, None, 350000]


def make_xlsx(tmp_path, with_preamble_rows=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit R-1"
    if with_preamble_rows:
        ws.append(["Exhibit R-1, RDT&E Programs"])  # banner rows above the header
        ws.append([])
    ws.append(HEADERS)
    ws.append(ROW1)
    ws.append(ROW2)
    p = tmp_path / "r1_display.xlsx"
    wb.save(p)
    return p


def test_load_rollup_melts_fy_columns(pg_dsn, tmp_path, doc_id):
    n = load_rollup(pg_dsn, make_xlsx(tmp_path), exhibit="R-1", fiscal_year=2026,
                    source_document_id=doc_id)
    assert n == 7  # ROW1: 4 amounts, ROW2: 3 amounts (None skipped)
    with psycopg.connect(pg_dsn) as con:
        val = con.execute(
            "select amount_thousands from budget_lines "
            "where pe_bli='0601101E' and amount_type='fy_2024_actuals'"
        ).fetchone()[0]
        assert val == Decimal("280494")
        types = {r[0] for r in con.execute(
            "select distinct amount_type from budget_lines where pe_bli='0602303E'")}
    assert types == {"fy_2024_actuals", "fy_2025_enacted", "fy_2026_disc_request"}


def test_load_rollup_is_idempotent_upsert(pg_dsn, tmp_path, doc_id):
    p = make_xlsx(tmp_path)
    load_rollup(pg_dsn, p, exhibit="R-1", fiscal_year=2026, source_document_id=doc_id)
    n2 = load_rollup(pg_dsn, p, exhibit="R-1", fiscal_year=2026, source_document_id=doc_id)
    assert n2 == 7  # upserts same rows
    with psycopg.connect(pg_dsn) as con:
        total = con.execute("select count(*) from budget_lines").fetchone()[0]
    assert total == 7


def test_load_rollup_records_cell_provenance(pg_dsn, tmp_path, doc_id):
    xlsx = make_xlsx(tmp_path, with_preamble_rows=True)
    load_rollup(pg_dsn, xlsx, exhibit="R-1", fiscal_year=2026, source_document_id=doc_id)
    with psycopg.connect(pg_dsn) as con:
        row = con.execute(
            "select source_sheet, source_cells from budget_lines"
            " where amount_type = 'fy_2024_actuals'"
            " and pe_bli = '0601101E'"
        ).fetchone()
    assert row[0]                      # sheet name recorded
    assert row[1] == ["J4"]            # col J (0-based 9 → letter J), data row 4


def test_load_rollup_skips_nan_amount(pg_dsn, tmp_path, doc_id):
    """Cells whose value is the literal string 'NaN' produce Decimal('NaN') without raising,
    so they must be caught by an explicit is_nan() guard and skipped."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit R-1"
    ws.append(HEADERS)
    # amount cell is the string "NaN" — Decimal("NaN") does NOT raise; is_nan() is True
    ws.append(["0400", "RDT&E Defense-Wide", "DARPA", "01", "Basic Research", "2",
               "0601101E", "DEFENSE RESEARCH SCIENCES", "Y",
               "NaN", 293145, None, 0])
    p = tmp_path / "r1_nan.xlsx"
    wb.save(p)

    load_rollup(pg_dsn, p, exhibit="R-1", fiscal_year=2026, source_document_id=doc_id)
    with psycopg.connect(pg_dsn) as con:
        # The NaN cell (fy_2024_actuals) must NOT be inserted
        nan_row = con.execute(
            "select count(*) from budget_lines"
            " where pe_bli='0601101E' and amount_type='fy_2024_actuals'"
        ).fetchone()[0]
        # But the valid cells should still be loaded
        valid_row = con.execute(
            "select count(*) from budget_lines"
            " where pe_bli='0601101E' and amount_type='fy_2025_enacted'"
        ).fetchone()[0]
    assert nan_row == 0, "NaN amount must be skipped (not inserted)"
    assert valid_row == 1, "Valid amounts in the same row must still load"


def test_split_ba_programs_keep_both_rows(pg_dsn, tmp_path, doc_id):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit R-1"
    ws.append(HEADERS)
    ws.append(["0400", "RDT&E AF", "F", "03", "Advanced Dev", "104",
               "0604776F", "DEPLOYMENT & DISTRIBUTION", "Y",
               19441, 19441, None, 20000])
    ws.append(["0400", "RDT&E AF", "F", "04", "Adv Component Dev", "131",
               "0604776F", "DEPLOYMENT & DISTRIBUTION", "Y",
               4840, 4840, None, 5000])
    p = tmp_path / "r1_split.xlsx"
    wb.save(p)

    load_rollup(pg_dsn, p, exhibit="R-1", fiscal_year=2026, source_document_id=doc_id)
    with psycopg.connect(pg_dsn) as con:
        rows = con.execute(
            "select budget_activity, amount_thousands from budget_lines "
            "where pe_bli='0604776F' and amount_type='fy_2025_enacted' order by 1"
        ).fetchall()
        total = con.execute(
            "select sum(amount_thousands) from budget_lines "
            "where pe_bli='0604776F' and amount_type='fy_2025_enacted'"
        ).fetchone()[0]
    assert len(rows) == 2
    assert total == Decimal("24281")


# --- PB2017–PB2023 era header variants (live workbook evidence, Task 4) ---

ERA_HEADERS_2017 = [
    "Account", "Account Title", "Organization", "Budget\nActivity",
    "Budget Activity Title", "Line\nNumber", "PE / BLI",
    "Program Element / Budget Line Item (BLI) Title", "Include\nin\nTOA",
    "FY 2015\n(Base & OCO)", "FY 2016\nTotal Enacted", "FY 2017\nBase",
    "FY 2017\nTotal", "Classification",
]


def make_era_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit R-1"
    ws.append(["Total of Displayed Rows"])
    ws.append(ERA_HEADERS_2017)
    ws.append(["0400", "RDT&E Defense-Wide", "DTRA", "02", "Applied Research",
               "18", "0602718BR", "WMD DEFEAT TECHNOLOGIES", "Y",
               172869, 165786, 160000, 163340, "U"])
    p = tmp_path / "r1_display.xlsx"
    wb.save(p)
    return p


def test_load_rollup_normalizes_era_headers(pg_dsn, tmp_path, doc_id):
    """PB2017–PB2023 display workbooks wrap headers across lines and space
    the slash in 'PE / BLI'; header matching normalizes whitespace and
    slashes so the era loads without touching modern behavior."""
    n = load_rollup(pg_dsn, make_era_xlsx(tmp_path), exhibit="R-1", fiscal_year=2017,
                    source_document_id=doc_id)
    assert n == 4
    with psycopg.connect(pg_dsn) as con:
        rows = dict(con.execute(
            "select amount_type, amount_thousands from budget_lines"
            " where pe_bli='0602718BR'"
        ).fetchall())
        org, ba = con.execute(
            "select organization, budget_activity from budget_lines"
            " where pe_bli='0602718BR' limit 1"
        ).fetchone()
    assert rows == {
        "fy_2015_base_oco": Decimal("172869"),
        "fy_2016_total_enacted": Decimal("165786"),
        "fy_2017_base": Decimal("160000"),
        "fy_2017_total": Decimal("163340"),
    }
    assert org == "DTRA"
    assert ba == "02"  # 'Budget\\nActivity' resolves to the budget_activity id
